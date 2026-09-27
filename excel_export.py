# -*- coding: utf-8 -*-
"""Экспорт сводки судебных дел в Excel (трек №2: результат → Excel юристу).

Использование:
  python excel_export.py --input СВОДКА_результатов.json --output документы_судов.xlsx
"""
import argparse
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "Номер дела", "Тип документа", "Суд", "Адрес суда",
    "Взыскатель", "ИНН взыскателя", "Адрес взыскателя",
    "Должник", "ИНН должника", "Адрес должника",
    "Основной долг", "Неустойка", "Госпошлина", "Итого",
    "Источник файл",
]


def row_of(x):
    c = x.get("court", {})
    cr = x.get("creditor", {})
    d = x.get("debtor", {})
    a = x.get("amounts", {})
    return [
        x.get("case_number", ""), x.get("document_type", ""),
        c.get("name", ""), c.get("address", ""),
        cr.get("name", ""), cr.get("inn", ""), cr.get("address", ""),
        d.get("name", ""), d.get("inn", ""), d.get("address", ""),
        a.get("principal"), a.get("penalty"), a.get("state_duty"), a.get("total"),
        x.get("_source", ""),
    ]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="путь к JSON-сводке (*_result.json списком)")
    ap.add_argument("--output", required=True, help="путь к выходному .xlsx")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, dict) and "cases" in data:  # совместимость с обёрткой
        data = data["cases"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Судебные акты"
    ws.append(HEADERS)
    hfill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hfill
    n_rows = 0
    for x in data:
        if not x.get("case_number"):
            continue
        ws.append(row_of(x))
        n_rows += 1
    widths = [18, 14, 34, 40, 38, 14, 40, 38, 14, 40, 12, 12, 12, 12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(args.output)
    print("Excel сохранён:", args.output)
    print("Строк:", n_rows)


if __name__ == "__main__":
    main()